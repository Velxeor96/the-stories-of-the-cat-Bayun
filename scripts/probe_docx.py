# PROBE_DOCX_V1 — диагностика структуры .docx
from __future__ import annotations
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parent.parent
BOOK = ROOT / "OKP_Rogue_Trader_rus_v1_23_clean_3.docx"

if not BOOK.exists():
    print("[ERROR] Нет файла:", BOOK)
    raise SystemExit(1)

doc = Document(str(BOOK))
print("[info] Всего параграфов:", len(doc.paragraphs))
print("[info] Всего таблиц:", len(doc.tables))
print("")

print("=== первые 120 непустых параграфов с указанием стиля ===")
n = 0
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if not text:
        continue
    style = p.style.name if p.style else "?"
    # укорачиваем для читаемости
    short = text[:110].replace("\n", " ")
    print(f"{i:05d} | {style:<20} | {short}")
    n += 1
    if n >= 120:
        break

print("")
print("=== Поиск параграфов со словом 'Глава' ===")
found = 0
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if "Глава" in text and len(text) < 120:
        style = p.style.name if p.style else "?"
        print(f"{i:05d} | {style:<20} | {text}")
        found += 1
        if found >= 40:
            break
print("[info] найдено глав-параграфов:", found)

print("")
print("=== Уникальные стили ===")
styles = set()
for p in doc.paragraphs:
    if p.text.strip() and p.style:
        styles.add(p.style.name)
for s in sorted(styles):
    print("  ", s)
